import csv
import io
from datetime import datetime
from flask import Response, make_response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def export_personnel_csv(personnel_data):
    """
    Génère un export CSV encodé en UTF-8-SIG (compatible nativement avec Excel Windows)
    """
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    
    # En-têtes
    writer.writerow(["Type", "Nom", "Prenom", "Email", "Specialite / Role", "Statut Activation", "Expiration Token"])
    
    for p in personnel_data:
        writer.writerow([
            p.get('type', ''),
            p.get('nom', ''),
            p.get('prenom', ''),
            p.get('email', ''),
            p.get('specialite', '-'),
            p.get('statut', ''),
            p.get('date_expiration', '-')
        ])
    
    # Conversion en bytes UTF-8 avec BOM (utf-8-sig)
    csv_bytes = output.getvalue().encode('utf-8-sig')
    filename = f"personnel_sante_scolaire_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    
    response = Response(csv_bytes, mimetype='text/csv; charset=utf-8')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_personnel_excel(personnel_data, kpis_data=None):
    """
    Génère un fichier Excel (.xlsx) professionnel avec mise en forme openpyxl.
    """
    wb = Workbook()
    
    # --- Feuille 1: Effectif & Personnel ---
    ws1 = wb.active
    ws1.title = "Effectif Médical"
    
    # Couleurs du thème médical
    primary_color = "1E3A8A"  # Bleu foncé
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Titre du rapport
    ws1.merge_cells("A1:G1")
    ws1["A1"] = f"RÉPERTOIRE DU PERSONNEL DE SANTÉ SCOLAIRE — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[1].height = 30
    
    # En-têtes de colonnes
    headers = ["Type", "Nom", "Prénom", "Adresse Email", "Spécialité / Fonction", "Statut Compte", "Expiration Token"]
    ws1.append([]) # Ligne vide 2
    ws1.append(headers) # Ligne 3
    ws1.row_dimensions[3].height = 24
    
    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    for row_idx, p in enumerate(personnel_data, start=4):
        ws1.row_dimensions[row_idx].height = 20
        row_values = [
            p.get('type', ''),
            p.get('nom', ''),
            p.get('prenom', ''),
            p.get('email', ''),
            p.get('specialite', '-'),
            p.get('statut', ''),
            p.get('date_expiration', '-')
        ]
        ws1.append(row_values)
        
        # Coloration zébrée et bordures
        for col_num in range(1, len(row_values) + 1):
            cell = ws1.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if row_idx % 2 == 0:
                cell.fill = zebra_fill
                
            if col_num in [1, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            # Couleur statut
            if col_num == 6:
                if p.get('statut') == 'Actif':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="15803D") # Vert
                elif p.get('statut') == 'Expiré':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C") # Rouge
                else:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B45309") # Ambre
    
    # Auto-dimensionnement des colonnes
    from openpyxl.utils import get_column_letter
    for col_idx, col in enumerate(ws1.columns, start=1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # --- Feuille 2: Synthèse Exécutive (KPIs) ---
    if kpis_data:
        ws2 = wb.create_sheet(title="Synthèse & Indicateurs")
        ws2.merge_cells("A1:C1")
        ws2["A1"] = "SYNTHÈSE DU PILOTAGE DE L'ÉTABLISSEMENT"
        ws2["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        ws2.row_dimensions[1].height = 30
        
        ws2.append([])
        ws2.append(["Indicateur Clé (KPI)", "Valeur Mesurée", "Interprétation"])
        ws2.row_dimensions[3].height = 24
        
        for col_num in range(1, 4):
            cell = ws2.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
        kpi_rows = [
            ("Personnel de santé total", kpis_data.get('total_personnel', 0), "Effectif médical et paramédical actif ou invité"),
            ("Médecins scolaires", kpis_data.get('total_medecins', 0), "Médecins affectés à l'établissement"),
            ("Infirmiers scolaires", kpis_data.get('total_infirmiers', 0), "Personnel soignant présent"),
            ("Taux d'activation des comptes (48h)", f"{kpis_data.get('taux_activation', 0)} %", "Ratio de comptes activés avec succès"),
            ("Nombre de spécialités médicales", kpis_data.get('nb_specialites', 0), "Spécialités couvertes"),
            ("Élèves inscrits", kpis_data.get('total_eleves', 0), "Effectif total des apprenants"),
            ("Élèves avec suivi PAI / Antécédents", kpis_data.get('eleves_avec_pai', 0), "Cas chroniques ou allergies sévères"),
            ("Consultations réalisées", kpis_data.get('total_consultations', 0), "Actes médicaux enregistrés"),
            ("Rendez-vous programmés", kpis_data.get('total_rdv', 0), "Consultations futures ou planifiées"),
            ("Médicaments en stock critique", kpis_data.get('stocks_critiques', 0), "Articles sous le seuil de réapprovisionnement")
        ]
        
        for row_idx, (nom_kpi, val_kpi, comment) in enumerate(kpi_rows, start=4):
            ws2.row_dimensions[row_idx].height = 22
            ws2.append([nom_kpi, val_kpi, comment])
            for col_num in range(1, 4):
                cell = ws2.cell(row=row_idx, column=col_num)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if row_idx % 2 == 0:
                    cell.fill = zebra_fill
                if col_num == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = Font(name="Calibri", size=10, bold=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    
        for col_idx, col in enumerate(ws2.columns, start=1):
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws2.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 4, 15)

        # Bloc de visa & signature dans Excel
        sig_row = len(kpi_rows) + 6
        ws2.cell(row=sig_row, column=1, value="VISAS & SIGNATURES D'APPROBATION").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        ws2.cell(row=sig_row + 1, column=1, value="Responsable Médical : _______________________").font = Font(name="Calibri", size=10, italic=True)
        ws2.cell(row=sig_row + 1, column=3, value="Direction d'Établissement : _______________________").font = Font(name="Calibri", size=10, italic=True)
        ws2.cell(row=sig_row + 2, column=1, value="Date & Signature : _____ / _____ / 202___").font = Font(name="Calibri", size=9, color="64748B")
        ws2.cell(row=sig_row + 2, column=3, value="Cachet Officiel : _____ / _____ / 202___").font = Font(name="Calibri", size=9, color="64748B")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"rapport_pilotage_santescolaire_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = Response(buffer.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_pilotage_pdf(personnel_data, kpis_data, alertes_data):
    """
    Génère un rapport exécutif de pilotage au format PDF (ReportLab)
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=0,
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )
    
    section_style = ParagraphStyle(
        'DocSection',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=8
    )
    
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11
    )
    
    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )
    
    # 1. En-tête
    elements.append(Paragraph("SANTÉ SCOLAIRE — TABLEAU DE BORD DE PILOTAGE", title_style))
    elements.append(Paragraph(f"Rapport officiel généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')} | Direction des Services Médicaux", subtitle_style))
    elements.append(Spacer(1, 10))
    
    # 2. Table des Indicateurs Clés (KPIs)
    elements.append(Paragraph("1. Indicateurs Clés de Performance (KPIs)", section_style))
    kpi_table_data = [
        [
            Paragraph("Indicateur", cell_header),
            Paragraph("Valeur", cell_header),
            Paragraph("Indicateur", cell_header),
            Paragraph("Valeur", cell_header)
        ],
        [
            Paragraph("Personnel total", cell_style),
            Paragraph(str(kpis_data.get('total_personnel', 0)), cell_style),
            Paragraph("Élèves inscrits", cell_style),
            Paragraph(str(kpis_data.get('total_eleves', 0)), cell_style)
        ],
        [
            Paragraph("Médecins en service", cell_style),
            Paragraph(str(kpis_data.get('total_medecins', 0)), cell_style),
            Paragraph("Consultations menées", cell_style),
            Paragraph(str(kpis_data.get('total_consultations', 0)), cell_style)
        ],
        [
            Paragraph("Infirmiers en poste", cell_style),
            Paragraph(str(kpis_data.get('total_infirmiers', 0)), cell_style),
            Paragraph("Rendez-vous programmés", cell_style),
            Paragraph(str(kpis_data.get('total_rdv', 0)), cell_style)
        ],
        [
            Paragraph("Taux d'activation (48h)", cell_style),
            Paragraph(f"{kpis_data.get('taux_activation', 0)} %", cell_style),
            Paragraph("Suivis PAI & Antécédents", cell_style),
            Paragraph(str(kpis_data.get('eleves_avec_pai', 0)), cell_style)
        ],
        [
            Paragraph("Spécialités couvertes", cell_style),
            Paragraph(str(kpis_data.get('nb_specialites', 0)), cell_style),
            Paragraph("Stocks critiques pharmacie", cell_style),
            Paragraph(str(kpis_data.get('stocks_critiques', 0)), cell_style)
        ]
    ]
    
    t_kpi = Table(kpi_table_data, colWidths=[160, 100, 160, 100])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_kpi)
    elements.append(Spacer(1, 10))
    
    # 3. Alertes et Points d'attention
    if alertes_data:
        elements.append(Paragraph(f"2. Points d'attention & Alertes Décisionnelles ({len(alertes_data)})", section_style))
        alert_table_data = [[
            Paragraph("Niveau", cell_header),
            Paragraph("Alerte", cell_header),
            Paragraph("Détails / Recommandation", cell_header)
        ]]
        for alt in alertes_data[:8]: # Max 8 alertes sur le résumé PDF
            alert_table_data.append([
                Paragraph(alt.get('type', '').upper(), cell_style),
                Paragraph(alt.get('titre', ''), cell_style),
                Paragraph(alt.get('message', ''), cell_style)
            ])
            
        t_alert = Table(alert_table_data, colWidths=[70, 150, 300])
        t_alert.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DC2626')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FEF2F2')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_alert)
        elements.append(Spacer(1, 10))
        
    # 4. Effectif Médical (Échantillon)
    elements.append(Paragraph(f"3. Répertoire du Personnel Médical ({len(personnel_data)})", section_style))
    perso_table_data = [[
        Paragraph("Type", cell_header),
        Paragraph("Nom & Prénom", cell_header),
        Paragraph("Email Professionnel", cell_header),
        Paragraph("Spécialité", cell_header),
        Paragraph("Statut", cell_header)
    ]]
    for p in personnel_data:
        perso_table_data.append([
            Paragraph(p.get('type', ''), cell_style),
            Paragraph(f"{p.get('prenom', '')} {p.get('nom', '')}", cell_style),
            Paragraph(p.get('email', ''), cell_style),
            Paragraph(p.get('specialite', '-'), cell_style),
            Paragraph(p.get('statut', ''), cell_style)
        ])
        
    t_perso = Table(perso_table_data, colWidths=[65, 120, 175, 100, 60])
    t_perso.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(t_perso)
    elements.append(Spacer(1, 18))
    
    # 5. Page / Section de Validation & Signatures Officielles
    elements.append(Paragraph("4. Émargement, Approbation & Visas Légaux", section_style))
    certif_text = Paragraph(
        "<i>« Je soussigné certifie sur l'honneur l'exactitude des données de pilotage et d'activité clinique ci-dessus présentées, "
        "établies conformément aux exigences de traçabilité médico-légale et de secret médical de la santé scolaire. »</i>",
        cell_style
    )
    elements.append(certif_text)
    elements.append(Spacer(1, 12))

    signature_headers = ParagraphStyle(
        'SigHeader',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        fontName='Helvetica-Bold',
        textColor=colors.HexColor('#1E3A8A')
    )
    signature_instructions = ParagraphStyle(
        'SigInst',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#64748B')
    )

    signatures_data = [
        [
            Paragraph("<b>POUR LE CORPS MÉDICAL</b><br/>Médecin Référent / Chef de Service", signature_headers),
            Paragraph("<b>POUR LA DIRECTION DE L'ÉTABLISSEMENT</b><br/>Chef d'Établissement / Proviseur", signature_headers)
        ],
        [
            Paragraph(
                "Date : _____ / _____ / 202___<br/><br/>"
                "Nom &amp; Prénom : ______________________<br/><br/>"
                "Signature &amp; Cachet Médical :<br/><br/><br/><br/>",
                signature_instructions
            ),
            Paragraph(
                "Date : _____ / _____ / 202___<br/><br/>"
                "Nom &amp; Prénom : ______________________<br/><br/>"
                "Signature &amp; Sceau Officiel de l'Établissement :<br/><br/><br/><br/>",
                signature_instructions
            )
        ]
    ]

    t_signatures = Table(signatures_data, colWidths=[255, 255])
    t_signatures.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94A3B8')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(t_signatures)
    elements.append(Spacer(1, 10))

    footer_legal = Paragraph(
        "<font size='7' color='#94A3B8'>Document généré numériquement via SanteScolaire Pro. "
        "Toute modification manuscrite non paraphée rend le présent extrait caduc. "
        "Conservation obligatoire au registre médical de l'établissement.</font>",
        ParagraphStyle('Foot', parent=styles['Normal'], alignment=1)
    )
    elements.append(footer_legal)
    
    doc.build(elements)
    buffer.seek(0)
    
    filename = f"rapport_pilotage_santescolaire_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    response = Response(buffer.getvalue(), mimetype='application/pdf')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
